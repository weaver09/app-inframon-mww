import os
import io
import csv
import json

from azure.identity import (
    DefaultAzureCredential,
    get_bearer_token_provider
)

from azure.storage.blob import BlobServiceClient
from openai import OpenAI
from mssql_python import connect
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


# ============================================================
# Configuration
# ============================================================

SQL_CONNECTION_STRING = os.getenv(
    "AZURE_SQL_CONNECTIONSTRING"
)

STORAGE_ACCOUNT_NAME = os.getenv(
    "AZURE_STORAGE_ACCOUNT"
)

AZURE_OPENAI_ENDPOINT = os.getenv(
    "AZURE_OPENAI_ENDPOINT"
)

AZURE_OPENAI_DEPLOYMENT = os.getenv(
    "AZURE_OPENAI_DEPLOYMENT"
)

STORAGE_CONTAINER_NAME = "documents"


ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "xlsx",
    "pptx",
    "txt",
    "csv",
    "json",
    "xml",
    "html",
    "htm",
    "md",
    "log",
    "ps1",
    "py",
    "sql",
    "yml",
    "yaml",
    "ini",
    "config"
}


TEXT_EXTENSIONS = {
    "txt",
    "json",
    "xml",
    "html",
    "htm",
    "md",
    "log",
    "ps1",
    "py",
    "sql",
    "yml",
    "yaml",
    "ini",
    "config"
}


# ============================================================
# Database
# ============================================================

def get_db_connection():

    if not SQL_CONNECTION_STRING:
        raise RuntimeError(
            "AZURE_SQL_CONNECTIONSTRING is not configured."
        )

    return connect(
        SQL_CONNECTION_STRING
    )


# ============================================================
# Azure Credential
# ============================================================

def get_azure_credential():

    return DefaultAzureCredential()


# ============================================================
# Blob Storage
# ============================================================

def get_blob_service_client():

    if not STORAGE_ACCOUNT_NAME:
        raise RuntimeError(
            "AZURE_STORAGE_ACCOUNT is not configured."
        )

    account_url = (
        f"https://{STORAGE_ACCOUNT_NAME}.blob.core.windows.net"
    )

    return BlobServiceClient(
        account_url=account_url,
        credential=get_azure_credential()
    )


# ============================================================
# Foundry / OpenAI
# ============================================================

def get_openai_client():

    if not AZURE_OPENAI_ENDPOINT:
        raise RuntimeError(
            "AZURE_OPENAI_ENDPOINT is not configured."
        )

    if not AZURE_OPENAI_DEPLOYMENT:
        raise RuntimeError(
            "AZURE_OPENAI_DEPLOYMENT is not configured."
        )

    credential = DefaultAzureCredential()

    token_provider = get_bearer_token_provider(
        credential,
        "https://cognitiveservices.azure.com/.default"
    )

    return OpenAI(
        base_url=AZURE_OPENAI_ENDPOINT,
        api_key=token_provider
    )


# ============================================================
# Helpers
# ============================================================

def allowed_file(filename):

    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower()
        in ALLOWED_EXTENSIONS
    )


def safe_json_loads(value):

    if not value:
        return []

    try:
        return json.loads(value)

    except Exception:
        return []


def decode_text_file(file_data):

    encodings = [
        "utf-8-sig",
        "utf-8",
        "utf-16",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:

        try:
            return file_data.decode(
                encoding
            )

        except UnicodeDecodeError:
            continue

    raise RuntimeError(
        "Unable to decode text file."
    )


# ============================================================
# PDF
# ============================================================

def extract_pdf(file_data):

    reader = PdfReader(
        io.BytesIO(file_data)
    )

    parts = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text() or ""

        if text.strip():

            parts.append(
                f"\n--- Page {page_number} ---\n{text}"
            )

    return "\n".join(parts)


# ============================================================
# DOCX
# ============================================================

def extract_docx(file_data):

    document = Document(
        io.BytesIO(file_data)
    )

    parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        parts.append(
            f"\n--- Table {table_index} ---"
        )

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            parts.append(
                " | ".join(values)
            )

    return "\n".join(parts)


# ============================================================
# XLSX
# ============================================================

def extract_xlsx(file_data):

    workbook = load_workbook(
        io.BytesIO(file_data),
        read_only=True,
        data_only=True
    )

    parts = []

    for worksheet in workbook.worksheets:

        parts.append(
            f"\n--- Worksheet: {worksheet.title} ---"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is None:
                    values.append("")

                else:
                    values.append(
                        str(value)
                    )

            if any(
                value.strip()
                for value in values
            ):

                parts.append(
                    " | ".join(values)
                )

    workbook.close()

    return "\n".join(parts)


# ============================================================
# PPTX
# ============================================================

def extract_pptx(file_data):

    presentation = Presentation(
        io.BytesIO(file_data)
    )

    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        parts.append(
            f"\n--- Slide {slide_number} ---"
        )

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):

                text = (
                    shape.text
                    .strip()
                )

                if text:
                    parts.append(text)

            if (
                hasattr(shape, "has_table")
                and shape.has_table
            ):

                for row in shape.table.rows:

                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    parts.append(
                        " | ".join(values)
                    )

    return "\n".join(parts)


# ============================================================
# CSV
# ============================================================

def extract_csv(file_data):

    text = decode_text_file(
        file_data
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    parts = []

    for row in reader:

        parts.append(
            " | ".join(
                str(value)
                for value in row
            )
        )

    return "\n".join(parts)


# ============================================================
# Text
# ============================================================

def extract_text_file(file_data):

    return decode_text_file(
        file_data
    )


# ============================================================
# Extract Content
# ============================================================

def extract_content(
    filename,
    file_data
):

    extension = (
        filename
        .lower()
        .rsplit(".", 1)[-1]
    )

    if extension == "pdf":
        return extract_pdf(file_data)

    if extension == "docx":
        return extract_docx(file_data)

    if extension == "xlsx":
        return extract_xlsx(file_data)

    if extension == "pptx":
        return extract_pptx(file_data)

    if extension == "csv":
        return extract_csv(file_data)

    if extension in TEXT_EXTENSIONS:
        return extract_text_file(
            file_data
        )

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


# ============================================================
# AI Analysis
# ============================================================

def analyze_document_with_ai(
    extracted_text
):

    if not extracted_text:

        raise RuntimeError(
            "No extracted text was available for AI analysis."
        )

    client = get_openai_client()

    max_characters = 60000

    text_to_analyze = (
        extracted_text[:max_characters]
    )

    system_message = """
You are an AI document analysis system.

Analyze documents accurately and conservatively.

Do not invent information.

Return only valid JSON.

If information is not present in the document,
use an empty array or null value where appropriate.
"""

    user_message = f"""
Analyze the following document.

Return ONLY valid JSON using exactly this structure:

{{
    "summary": "Concise but useful document summary",

    "key_topics": [
        "topic"
    ],

    "entities": [
        {{
            "name": "entity name",
            "type": "person|organization|product|location|system|other"
        }}
    ],

    "tags": [
        "tag"
    ],

    "risks": [
        {{
            "risk": "risk description",
            "severity": "Low|Medium|High"
        }}
    ],

    "action_items": [
        {{
            "action": "action description",
            "owner": null,
            "due_date": null
        }}
    ],

    "important_dates": [
        {{
            "date": "YYYY-MM-DD or original date text",
            "description": "why this date matters"
        }}
    ]
}}

Rules:

1. Do not invent facts.
2. Summary should be factual and useful.
3. Key topics should identify major themes.
4. Entities should identify important people,
   organizations, systems, products, and locations.
5. Tags should be short and useful for searching.
6. Risks must be supported by the document.
7. If no risks exist, return an empty array.
8. Action items must come from explicit or strongly
   implied tasks in the document.
9. Preserve owners when explicitly stated.
10. Preserve due dates when explicitly stated.
11. If no action items exist, return an empty array.
12. Important dates should only include meaningful dates.
13. If no important dates exist, return an empty array.

DOCUMENT:

{text_to_analyze}
"""

    response = client.responses.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        input=[
            {
                "role": "system",
                "content": system_message
            },
            {
                "role": "user",
                "content": user_message
            }
        ]
    )

    content = response.output_text

    if not content:

        raise RuntimeError(
            "Azure OpenAI returned an empty response."
        )

    try:
        analysis = json.loads(
            content
        )

    except json.JSONDecodeError as exc:

        raise RuntimeError(
            "Azure OpenAI returned invalid JSON."
        ) from exc

    analysis.setdefault(
        "summary",
        ""
    )

    analysis.setdefault(
        "key_topics",
        []
    )

    analysis.setdefault(
        "entities",
        []
    )

    analysis.setdefault(
        "tags",
        []
    )

    analysis.setdefault(
        "risks",
        []
    )

    analysis.setdefault(
        "action_items",
        []
    )

    analysis.setdefault(
        "important_dates",
        []
    )

    return analysis


# ============================================================
# Ask Document
# ============================================================

def ask_document_with_ai(
    extracted_text,
    question
):

    if not extracted_text:
        raise RuntimeError(
            "No extracted text is available."
        )

    if not question:
        raise RuntimeError(
            "A question is required."
        )

    client = get_openai_client()

    text_to_analyze = (
        extracted_text[:60000]
    )

    system_message = """
You answer questions about a supplied document.

Only use information contained in the document.

Do not invent facts.

If the answer cannot be determined from the document,
say that clearly.

Provide a concise but useful answer.
"""

    user_message = f"""
DOCUMENT:

{text_to_analyze}


QUESTION:

{question}
"""

    response = client.responses.create(
        model=AZURE_OPENAI_DEPLOYMENT,
        input=[
            {
                "role": "system",
                "content": system_message
            },
            {
                "role": "user",
                "content": user_message
            }
        ]
    )

    answer = response.output_text

    if not answer:

        raise RuntimeError(
            "Azure OpenAI returned an empty response."
        )

    return answer


# ============================================================
# Save Analysis
# ============================================================

def save_analysis(
    cursor,
    document_id,
    analysis
):

    cursor.execute(
        """
        SELECT AnalysisID
        FROM dbo.DocumentAnalysis
        WHERE DocumentID = ?;
        """,
        (
            document_id,
        )
    )

    existing = cursor.fetchone()

    values = (
        analysis.get("summary"),

        json.dumps(
            analysis.get(
                "key_topics",
                []
            )
        ),

        json.dumps(
            analysis.get(
                "entities",
                []
            )
        ),

        json.dumps(
            analysis.get(
                "tags",
                []
            )
        ),

        json.dumps(
            analysis.get(
                "risks",
                []
            )
        ),

        json.dumps(
            analysis.get(
                "action_items",
                []
            )
        ),

        json.dumps(
            analysis.get(
                "important_dates",
                []
            )
        )
    )

    if existing:

        cursor.execute(
            """
            UPDATE dbo.DocumentAnalysis
            SET
                Summary = ?,
                KeyTopics = ?,
                EntitiesJson = ?,
                TagsJson = ?,
                RisksJson = ?,
                ActionItemsJson = ?,
                ImportantDatesJson = ?,
                AnalysisDate = SYSUTCDATETIME()
            WHERE DocumentID = ?;
            """,
            values + (
                document_id,
            )
        )

    else:

        cursor.execute(
            """
            INSERT INTO dbo.DocumentAnalysis
            (
                DocumentID,
                Summary,
                KeyTopics,
                EntitiesJson,
                TagsJson,
                RisksJson,
                ActionItemsJson,
                ImportantDatesJson
            )
            VALUES
            (
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?
            );
            """,
            (
                document_id,
            ) + values
        )